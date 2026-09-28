#!/usr/bin/env python3
"""
Add a register to CCGX-Modbus-TCP-register-list.xlsx.

Inserts a row on the "Field list" sheet in address order within the
service's block, copying the formatting of the neighbouring row. If the new
register falls inside a RESERVED row of the same service, that row is shrunk
or split. A one-line summary is appended to the "Document versions" sheet
(without a Rev number, those are added manually on release).

Requires openpyxl (pip install openpyxl).
"""
import argparse
import re
import xml.etree.ElementTree as ET
import zipfile
from copy import copy
from decimal import Decimal

try:
    import openpyxl
except ImportError:
    raise SystemExit("### openpyxl is required: pip install openpyxl")

FIELD_SHEET = 'Field list'
VERSIONS_SHEET = 'Document versions'
FIRST_DATA_ROW = 3

# Column numbers on the Field list sheet
COL_SERVICE = 1
COL_DESCRIPTION = 2
COL_ADDRESS = 3
COL_TYPE = 4
COL_SCALE = 5
COL_RANGE = 6
COL_PATH = 7
COL_WRITABLE = 8
COL_UNIT = 9
COL_REMARKS = 10
LAST_COL = COL_REMARKS

TYPE_LIMITS = {
    'uint16': (0, 65535),
    'int16': (-32768, 32767),
    'uint32': (0, 4294967295),
    'int32': (-2147483648, 2147483647),
}

parser = argparse.ArgumentParser(description='Add a register to the Modbus TCP register list xlsx')
parser.add_argument('--xlsx', default='CCGX-Modbus-TCP-register-list.xlsx', help='Path to the xlsx file')
parser.add_argument('--service', required=True, help='e.g. com.victronenergy.vebus')
parser.add_argument('--description', required=True, help='Human readable description')
parser.add_argument('--address', required=True, type=int, help='Modbus register address')
parser.add_argument('--type', required=True, help='uint16, int16, uint32, int32 or string[N]')
parser.add_argument('--scale', required=True, help='Scale factor, as in attributes.csv')
parser.add_argument('--path', required=True, help='D-Bus path')
parser.add_argument('--writable', required=True, choices=['yes', 'no'])
parser.add_argument('--unit', default='', help='Unit or enum text, e.g. W or 0=Off;1=On')
parser.add_argument('--remarks', default='', help='Optional remarks')
parser.add_argument('--range', dest='range_text', help='Override the computed Range column')
parser.add_argument('--doc-summary', help='Line to append to the Document versions sheet')
args = parser.parse_args()


def fail(msg):
    raise SystemExit(f"### {msg}")


def registerCount(modbusType):
    m = re.fullmatch(r'string\[(\d+)\]', modbusType)
    if m:
        return int(m.group(1))
    if modbusType.endswith('int32'):
        return 2
    return 1


def fmt(value):
    s = format(value.normalize(), 'f')
    return s.rstrip('0').rstrip('.') if '.' in s else s


def rangeText(modbusType, scale):
    m = re.fullmatch(r'string\[(\d+)\]', modbusType)
    if m:
        return f"{int(m.group(1)) * 2} characters"
    if modbusType not in TYPE_LIMITS:
        fail(f"unknown modbus type '{modbusType}'")
    lo, hi = TYPE_LIMITS[modbusType]
    s = Decimal(scale)
    if s == 0:
        fail("scale factor cannot be 0")
    return f"{fmt(Decimal(lo) / s)} to {fmt(Decimal(hi) / s)}"


def parseAddress(value):
    """Return (start, end) for an address cell, which is a number or 'a-b'."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value), int(value)
    m = re.fullmatch(r'\s*(\d+)\s*(?:-\s*(\d+))?\s*', str(value))
    if not m:
        return None
    start = int(m.group(1))
    return start, int(m.group(2) or start)


def addressValue(start, end):
    return start if start == end else f"{start}-{end}"


def numeric(value):
    """Store numbers as numbers, like the existing cells."""
    try:
        d = Decimal(value)
    except Exception:
        return value
    return int(d) if d == d.to_integral_value() else float(d)


def isReserved(ws, row):
    return str(ws.cell(row, COL_DESCRIPTION).value or '').strip().upper() == 'RESERVED'


def rowStyle(ws, row):
    return [copy(ws.cell(row, col)._style) for col in range(1, LAST_COL + 1)]


def insertRow(ws, idx, style):
    ws.insert_rows(idx)
    for col, s in enumerate(style, start=1):
        ws.cell(idx, col)._style = copy(s)


def fixNamedStyleIds(wb, xlsxPath):
    """
    openpyxl renumbers the named cell styles (cellStyleXfs) on save, dropping
    the unnamed ones, but leaves each cell format's xfId pointing at the
    original index. Excel then reports the file as corrupt. Remap the xfIds to
    the new positions of the named styles.
    """
    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(xlsxPath) as z:
        root = ET.fromstring(z.read('xl/styles.xml'))
    cellStyles = root.findall('m:cellStyles/m:cellStyle', ns)
    mapping = {int(cs.get('xfId')): i for i, cs in enumerate(cellStyles)}

    def remap(style):
        style.xfId = mapping.get(style.xfId, 0)

    for style in wb._cell_styles:
        remap(style)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.has_style:
                    remap(cell._style)


def main():
    wb = openpyxl.load_workbook(args.xlsx)
    fixNamedStyleIds(wb, args.xlsx)
    if FIELD_SHEET not in wb.sheetnames:
        fail(f"sheet '{FIELD_SHEET}' not found")
    ws = wb[FIELD_SHEET]

    newStart = args.address
    newEnd = args.address + registerCount(args.type) - 1

    serviceRows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if ws.cell(r, COL_SERVICE).value == args.service:
            addr = parseAddress(ws.cell(r, COL_ADDRESS).value)
            if addr:
                start, end = addr
                modbusType = str(ws.cell(r, COL_TYPE).value or '').strip().lower()
                if start == end and not isReserved(ws, r):
                    end = start + registerCount(modbusType) - 1
                serviceRows.append((r, (start, end)))

    reservedRow = None
    for r, (start, end) in serviceRows:
        if start <= newEnd and newStart <= end:
            if isReserved(ws, r) and start <= newStart and newEnd <= end:
                reservedRow = (r, start, end)
            else:
                fail(f"address {newStart}-{newEnd} overlaps row {r} ({start}-{end}) "
                     f"{ws.cell(r, COL_PATH).value or ws.cell(r, COL_DESCRIPTION).value}")

    if not serviceRows:
        fail(f"no rows for service {args.service} found; add the first row by hand")

    if reservedRow:
        idx = reservedRow[0]
    else:
        preceding = [r for r, (start, _) in serviceRows if start < newStart]
        idx = (max(preceding) + 1) if preceding else serviceRows[0][0]

    # Copy formatting from the nearest real register of this service
    candidates = [r for r, _ in serviceRows if not isReserved(ws, r)] or [r for r, _ in serviceRows]
    style = rowStyle(ws, min(candidates, key=lambda r: abs(r - idx)))

    if reservedRow:
        r, start, end = reservedRow
        before = (start, newStart - 1) if start < newStart else None
        after = (newEnd + 1, end) if newEnd < end else None
        if before and after:
            # Split: keep the first part in row r, add a row for the remainder
            ws.cell(r, COL_ADDRESS).value = addressValue(*before)
            reservedValues = [ws.cell(r, col).value for col in range(1, LAST_COL + 1)]
            insertRow(ws, r + 1, rowStyle(ws, r))
            for col, value in enumerate(reservedValues, start=1):
                ws.cell(r + 1, col).value = value
            ws.cell(r + 1, COL_ADDRESS).value = addressValue(*after)
            idx = r + 1
        elif before:
            ws.cell(r, COL_ADDRESS).value = addressValue(*before)
            idx = r + 1
        elif after:
            ws.cell(r, COL_ADDRESS).value = addressValue(*after)
            idx = r
        else:
            ws.delete_rows(r)
            idx = r
        print(f"Adjusted RESERVED {addressValue(start, end)}: "
              f"now {', '.join(str(addressValue(*p)) for p in (before, after) if p) or 'removed'}")

    insertRow(ws, idx, style)
    values = {
        COL_SERVICE: args.service,
        COL_DESCRIPTION: args.description,
        COL_ADDRESS: args.address,
        COL_TYPE: args.type,
        COL_SCALE: numeric(args.scale),
        COL_RANGE: args.range_text or rangeText(args.type, args.scale),
        COL_PATH: args.path,
        COL_WRITABLE: args.writable,
        COL_UNIT: args.unit or None,
        COL_REMARKS: args.remarks or None,
    }
    for col, value in values.items():
        ws.cell(idx, col).value = value
    print(f"Inserted row {idx}: " + ', '.join(str(v) for v in values.values() if v is not None))

    # Fix hyperlink references after rows moved
    for row in ws.iter_rows():
        for cell in row:
            if cell.hyperlink is not None:
                cell.hyperlink.ref = cell.coordinate

    if args.doc_summary:
        vs = wb[VERSIONS_SHEET]
        last = max((c.row for row in vs.iter_rows() for c in row if c.value not in (None, '')), default=0)
        vs.cell(last + 1, 2).value = args.doc_summary
        if vs.cell(last, 2).has_style:
            vs.cell(last + 1, 2)._style = copy(vs.cell(last, 2)._style)
        print(f"Added to '{VERSIONS_SHEET}' row {last + 1}: {args.doc_summary}")

    wb.save(args.xlsx)


if __name__ == '__main__':
    main()
